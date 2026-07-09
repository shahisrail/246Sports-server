#!/usr/bin/env python3
"""
Extract client-provided federation, calendar, and fixture files into reviewable
JSON/CSV data for the 246Sports importer.
"""

from __future__ import annotations

import csv
import json
import re
import subprocess
import sys
from datetime import datetime
from pathlib import Path
from typing import Any

import pdfplumber
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "tmp" / "client-data-import"
CLIENT_DATA_DIR = ROOT / "client-data"
SOURCES = {
    "federations": CLIENT_DATA_DIR / "National Sporting Federations- MSCE.xlsx",
    "sailing_xls": CLIENT_DATA_DIR / "Sailing Calendar 2026.xls",
    "sailing_xlsx": OUTPUT_DIR / "Sailing Calendar 2026.xlsx",
    "bca_dir": CLIENT_DATA_DIR / "BCA__Calendar_of_events_&_fixtures",
    "round_pdf": CLIENT_DATA_DIR / "2025 -2026 round 2 final by iScanner.pdf",
}

MONTHS = {
    "january": 1,
    "february": 2,
    "march": 3,
    "april": 4,
    "may": 5,
    "june": 6,
    "july": 7,
    "august": 8,
    "september": 9,
    "october": 10,
    "november": 11,
    "december": 12,
}

SPORT_ALIASES = {
    "Athletics (Track & Field)": "Track and Field",
    "Football (Soccer)": "Football",
    "Hockey (Field Hockey)": "Field Hockey",
    "Motorsports": "Motorsport",
}

SKIP_FIXTURE_WORDS = {
    "bye",
    "reserve day",
    "semi-finals",
    "semifinals",
    "final",
    "teams to be determined",
    "tbd",
    "...",
}


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\xa0", " ")
    return re.sub(r"\s+", " ", text).strip()


def slugify(value: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")
    return slug or "item"


def split_lines(value: Any) -> list[str]:
    if value is None:
        return []
    return [clean_text(part) for part in str(value).splitlines() if clean_text(part)]


def first_email(value: str) -> str:
    emails = re.findall(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", value or "")
    return emails[0] if emails else ""


def parse_month_date(raw: str, default_year: int = 2026) -> str:
    text = clean_text(raw)
    if not text:
        return ""

    match = re.search(
        r"(January|February|March|April|May|June|July|August|September|October|November|December)"
        r"\s+(\d{1,2})(?:st|nd|rd|th)?(?:,?\s*(\d{4}))?",
        text,
        re.IGNORECASE,
    )
    if match:
        month = MONTHS[match.group(1).lower()]
        day = int(match.group(2))
        year = int(match.group(3) or default_year)
        return datetime(year, month, day).date().isoformat()
    return ""


def parse_sailing_date(month: str, date_value: Any, default_year: int = 2026) -> tuple[str, str]:
    month = clean_text(month)
    date_text = clean_text(date_value)
    if not month or not date_text or month.lower() not in MONTHS:
        return "", ""

    days = [int(item) for item in re.findall(r"\d{1,2}", date_text)]
    if not days:
        return "", ""

    start = datetime(default_year, MONTHS[month.lower()], days[0]).date().isoformat()
    end = datetime(default_year, MONTHS[month.lower()], days[-1]).date().isoformat()
    return start, end


def competition_name_from_file(path: Path) -> str:
    stem = path.stem
    stem = re.sub(r"[_-]+", " ", stem)
    stem = re.sub(r"\b2\s+1\b", "", stem)
    stem = re.sub(r"\s+", " ", stem).strip()
    return stem.title().replace("Bbl", "BBL").replace("T20", "T20").replace("Bmf", "BMF")


def competition_category(name: str) -> str:
    upper = name.upper()
    if any(word in upper for word in ["CUP", "SHIELD", "PLATE"]):
        return "CUP"
    if any(word in upper for word in ["CALENDAR", "REGATTA", "RACING"]):
        return "TOURNAMENT"
    if "DIVISION" in upper:
        return "LEAGUE"
    return "OTHER"


def should_skip_name(value: str) -> bool:
    lower = clean_text(value).lower()
    return not lower or lower in SKIP_FIXTURE_WORDS or any(word == lower for word in SKIP_FIXTURE_WORDS)


def extract_federations() -> tuple[list[dict[str, Any]], list[str]]:
    warnings: list[str] = []
    rows: list[dict[str, Any]] = []
    workbook = load_workbook(SOURCES["federations"], read_only=True, data_only=True)
    sheet = workbook["National Sporting Federations-B"]

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if not row[0] or not row[2]:
            continue
        sport_name = SPORT_ALIASES.get(clean_text(row[7]), clean_text(row[7]))
        email = first_email(clean_text(row[6])) or first_email(clean_text(row[12])) or first_email(clean_text(row[16]))
        rows.append(
            {
                "source": SOURCES["federations"].name,
                "abbreviation": clean_text(row[1]),
                "name": clean_text(row[2]),
                "phone": clean_text(row[3]) or clean_text(row[10]),
                "address": clean_text(row[4]),
                "parish": clean_text(row[5]),
                "email": email,
                "sport": sport_name,
                "primaryContact": clean_text(row[8]),
                "primaryPosition": clean_text(row[9]),
                "secondaryContact": clean_text(row[13]),
                "secondaryPosition": clean_text(row[14]),
                "social": clean_text(row[17]) or clean_text(row[18]),
                "website": clean_text(row[19]),
                "status": clean_text(row[20]),
                "notes": clean_text(row[21]),
            }
        )

    return rows, warnings


def ensure_sailing_xlsx() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    if SOURCES["sailing_xlsx"].exists():
        return
    converted_fallback = ROOT / "tmp" / "data-inspect" / "Sailing Calendar 2026.xlsx"
    if converted_fallback.exists():
        SOURCES["sailing_xlsx"].write_bytes(converted_fallback.read_bytes())
        return
    try:
        subprocess.run(
            [
                "soffice",
                "--headless",
                "--convert-to",
                "xlsx",
                "--outdir",
                str(OUTPUT_DIR),
                str(SOURCES["sailing_xls"]),
            ],
            check=True,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
        )
    except Exception as exc:
        raise RuntimeError(f"Could not convert {SOURCES['sailing_xls'].name}: {exc}") from exc


def extract_sailing_events() -> tuple[list[dict[str, Any]], list[str]]:
    warnings: list[str] = []
    events: list[dict[str, Any]] = []
    ensure_sailing_xlsx()

    workbook = load_workbook(SOURCES["sailing_xlsx"], read_only=True, data_only=True)
    sheet = workbook["Sheet1"]
    current_month = ""

    for row in sheet.iter_rows(min_row=5, values_only=True):
        month = clean_text(row[1] if len(row) > 1 else "")
        if month and month.lower() in MONTHS:
            current_month = month
        event_name = clean_text(row[4] if len(row) > 4 else "")
        if not event_name:
            continue
        start_date, end_date = parse_sailing_date(current_month, row[3] if len(row) > 3 else "")
        if not start_date:
            warnings.append(f"Sailing date not parsed for event: {event_name}")
        events.append(
            {
                "source": SOURCES["sailing_xls"].name,
                "association": "Barbados Sailing Association",
                "sport": "Sailing",
                "name": event_name,
                "category": "TOURNAMENT",
                "startDate": start_date,
                "endDate": end_date or start_date,
                "venue": clean_text(row[7] if len(row) > 7 else ""),
                "organiser": clean_text(row[6] if len(row) > 6 else ""),
                "sponsor": clean_text(row[5] if len(row) > 5 else ""),
                "rawDate": clean_text(row[3] if len(row) > 3 else ""),
                "rawDay": clean_text(row[2] if len(row) > 2 else ""),
            }
        )

    return events, warnings


def cricket_association_for_file(path: Path) -> tuple[str, str]:
    name = path.name.lower()
    if "bridge" in name or "bbl" in name:
        return "Barbados Bridge Leaue", "Bridge"
    if "racing" in name or "motorsport" in name or "bmf" in name:
        return "Barbados Motoring Federation Inc.", "Motorsport"
    return "Barbados Cricket Association", "Cricket"


def extract_cricket_tables(path: Path) -> tuple[list[dict[str, Any]], list[str]]:
    fixtures: list[dict[str, Any]] = []
    warnings: list[str] = []
    competition_name = competition_name_from_file(path)
    association, sport = cricket_association_for_file(path)

    try:
        with pdfplumber.open(str(path)) as pdf:
            for page_index, page in enumerate(pdf.pages, start=1):
                tables = page.extract_tables() or []
                current_round = ""
                for table in tables:
                    for row in table:
                        cells = [clean_text(cell) for cell in row]
                        if not any(cells):
                            continue
                        if cells[0].lower().startswith("series"):
                            current_round = cells[0]
                            continue
                        if len(cells) < 5:
                            continue
                        date_raw, home_raw, marker, away_raw, venue_raw = cells[:5]
                        if not date_raw or "date" in date_raw.lower():
                            continue
                        if marker and marker.upper() != "V":
                            continue
                        homes = split_lines(row[1])
                        aways = split_lines(row[3])
                        venues = split_lines(row[4])
                        count = min(len(homes), len(aways))
                        if count == 0:
                            continue
                        date_iso = parse_month_date(date_raw)
                        if not date_iso:
                            warnings.append(f"{path.name}: date not parsed on page {page_index}: {date_raw}")
                        for index in range(count):
                            home = homes[index]
                            away = aways[index]
                            venue = venues[index] if index < len(venues) else (venues[-1] if venues else "")
                            if should_skip_name(home) or should_skip_name(away):
                                continue
                            fixtures.append(
                                {
                                    "source": path.name,
                                    "association": association,
                                    "sport": sport,
                                    "competition": competition_name,
                                    "category": competition_category(competition_name),
                                    "round": current_round,
                                    "date": date_iso,
                                    "rawDate": date_raw,
                                    "homeTeam": home,
                                    "awayTeam": away,
                                    "venue": venue,
                                    "page": page_index,
                                }
                            )
    except Exception as exc:
        warnings.append(f"{path.name}: {exc}")

    return fixtures, warnings


def extract_text_calendar(path: Path) -> tuple[list[dict[str, Any]], list[str]]:
    events: list[dict[str, Any]] = []
    warnings: list[str] = []
    association, sport = cricket_association_for_file(path)
    competition_name = competition_name_from_file(path)

    try:
        with pdfplumber.open(str(path)) as pdf:
            text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    except Exception as exc:
        return events, [f"{path.name}: {exc}"]

    for line in text.splitlines():
        line = clean_text(line)
        if not line:
            continue
        date_iso = parse_month_date(line)
        if not date_iso:
            continue
        description = re.sub(
            r"^(January|February|March|April|May|June|July|August|September|October|November|December)"
            r"\s+\d{1,2}(?:st|nd|rd|th)?(?:,?\s*\d{4})?\s*",
            "",
            line,
            flags=re.IGNORECASE,
        ).strip(" -")
        if not description:
            continue
        events.append(
            {
                "source": path.name,
                "association": association,
                "sport": sport,
                "name": description,
                "category": competition_category(competition_name),
                "startDate": date_iso,
                "endDate": date_iso,
                "venue": "",
                "organiser": "",
                "sponsor": "",
                "rawDate": line,
            }
        )
    return events, warnings


def extract_round_pdf() -> tuple[list[dict[str, Any]], list[str]]:
    path = SOURCES["round_pdf"]
    fixtures: list[dict[str, Any]] = []
    warnings: list[str] = []
    current_month = "June"

    try:
        with pdfplumber.open(str(path)) as pdf:
            for page_index, page in enumerate(pdf.pages, start=1):
                for table in page.extract_tables() or []:
                    current_date = ""
                    for row in table:
                        cells = [clean_text(cell) for cell in row]
                        if len(cells) < 6 or "DATE" in cells[0].upper():
                            continue
                        if cells[0].upper() in MONTHS:
                            current_month = cells[0].title()
                            continue
                        if cells[0]:
                            current_date = cells[0]
                        match_code = cells[2]
                        if "-" not in match_code:
                            continue
                        date_match = re.search(r"(\d{1,2})", current_date)
                        date_iso = ""
                        if date_match and current_month.lower() in MONTHS:
                            date_iso = datetime(2026, MONTHS[current_month.lower()], int(date_match.group(1))).date().isoformat()
                        teams = [clean_text(part) for part in match_code.split("-", 1)]
                        if len(teams) != 2 or not teams[0] or not teams[1]:
                            continue
                        fixtures.append(
                            {
                                "source": path.name,
                                "association": "Barbados Netball Association",
                                "sport": "Netball",
                                "competition": "2025-2026 Round 2",
                                "category": "LEAGUE",
                                "round": clean_text(cells[3]),
                                "date": date_iso,
                                "time": clean_text(cells[4]),
                                "rawDate": current_date,
                                "homeTeam": teams[0],
                                "awayTeam": teams[1],
                                "venue": clean_text(cells[5]),
                                "page": page_index,
                            }
                        )
    except Exception as exc:
        warnings.append(f"{path.name}: {exc}")

    return fixtures, warnings


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    if not rows:
        return
    keys: list[str] = []
    for row in rows:
        for key in row:
            if key not in keys:
                keys.append(key)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=keys)
        writer.writeheader()
        writer.writerows(rows)


def main() -> int:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    warnings: list[str] = []

    federations, federation_warnings = extract_federations()
    warnings.extend(federation_warnings)

    sailing_events, sailing_warnings = extract_sailing_events()
    warnings.extend(sailing_warnings)

    fixtures: list[dict[str, Any]] = []
    calendar_events = sailing_events[:]

    for pdf_path in sorted(SOURCES["bca_dir"].glob("*.pdf")):
        source_name = pdf_path.name.lower()
        if any(token in source_name for token in ["calendar", "racing", "motorsport", "bmf"]):
            events, event_warnings = extract_text_calendar(pdf_path)
            calendar_events.extend(events)
            warnings.extend(event_warnings)
            continue
        extracted, fixture_warnings = extract_cricket_tables(pdf_path)
        fixtures.extend(extracted)
        warnings.extend(fixture_warnings)

    round_fixtures, round_warnings = extract_round_pdf()
    fixtures.extend(round_fixtures)
    warnings.extend(round_warnings)

    sports = sorted(
        {
            clean_text(row["sport"])
            for row in federations
            if clean_text(row.get("sport"))
        }
        | {clean_text(row["sport"]) for row in fixtures if clean_text(row.get("sport"))}
        | {clean_text(row["sport"]) for row in calendar_events if clean_text(row.get("sport"))}
    )

    competitions: dict[tuple[str, str], dict[str, Any]] = {}
    for fixture in fixtures:
        key = (fixture["association"], fixture["competition"])
        current = competitions.setdefault(
            key,
            {
                "source": fixture["source"],
                "association": fixture["association"],
                "sport": fixture["sport"],
                "name": fixture["competition"],
                "category": fixture["category"],
                "status": "UPCOMING",
                "startDate": fixture["date"],
                "endDate": fixture["date"],
            },
        )
        if fixture["date"]:
            dates = [date for date in [current.get("startDate"), current.get("endDate"), fixture["date"]] if date]
            current["startDate"] = min(dates)
            current["endDate"] = max(dates)

    for event in calendar_events:
        competitions[(event["association"], event["name"])] = {
            "source": event["source"],
            "association": event["association"],
            "sport": event["sport"],
            "name": event["name"],
            "category": event["category"],
            "status": "UPCOMING",
            "startDate": event["startDate"],
            "endDate": event["endDate"],
            "venue": event["venue"],
            "description": " | ".join(
                part
                for part in [
                    f"Organiser: {event['organiser']}" if event.get("organiser") else "",
                    f"Sponsor: {event['sponsor']}" if event.get("sponsor") else "",
                    f"Raw date: {event['rawDate']}" if event.get("rawDate") else "",
                ]
                if part
            ),
        }

    data = {
        "generatedAt": datetime.utcnow().replace(microsecond=0).isoformat() + "Z",
        "sources": {key: str(value.relative_to(ROOT)) for key, value in SOURCES.items() if value.exists()},
        "sports": [{"name": sport, "slug": slugify(sport)} for sport in sports],
        "federations": federations,
        "competitions": sorted(competitions.values(), key=lambda item: (item["association"], item["name"])),
        "fixtures": fixtures,
        "calendarEvents": calendar_events,
        "warnings": warnings,
        "summary": {
            "sports": len(sports),
            "federations": len(federations),
            "competitions": len(competitions),
            "fixtures": len(fixtures),
            "calendarEvents": len(calendar_events),
            "warnings": len(warnings),
        },
    }

    json_path = OUTPUT_DIR / "client-data-extract.json"
    json_path.write_text(json.dumps(data, indent=2), encoding="utf-8")
    write_csv(OUTPUT_DIR / "federations.csv", federations)
    write_csv(OUTPUT_DIR / "competitions.csv", data["competitions"])
    write_csv(OUTPUT_DIR / "fixtures.csv", fixtures)
    write_csv(OUTPUT_DIR / "calendar-events.csv", calendar_events)

    print(json.dumps(data["summary"], indent=2))
    print(f"wrote {json_path.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
